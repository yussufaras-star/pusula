"""Banka ekstresi (CSV / XLSX / HTML-xls) ayrıştırma.

Garanti, KuveytTürk, Vakıf için API yok; internet şubesinden indirilen
dosya yüklenir. Türk bankalarının çıktıları başlıktan önce bilgi satırları
taşır: başlık satırı anahtar kelimelerle bulunur, sütun eşlemesi tahmin
edilir, kullanıcı düzeltebilir. external_id içerikten türetilir; aynı
dosya tekrar yüklenince aynı kimlikler çıkar, çift kayıt oluşmaz.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from typing import Union

from pydantic import BaseModel, model_validator

Cell = Union[str, int, float, Decimal, date, datetime, None]

HEADER_SCAN_ROWS = 30
MAX_FILE_BYTES = 3 * 1024 * 1024  # Vercel istek gövdesi 4,5 MB; base64 ile ~3 MB dosya
EXCEL_EPOCH = date(1899, 12, 30)
# Excel seri tarih aralığı: 1954 – 2119; dışı tutar sayılır, tarih değil.
EXCEL_SERIAL_MIN = 20000
EXCEL_SERIAL_MAX = 80000
OLE2_MAGIC = b"\xd0\xcf\x11\xe0"
ZIP_MAGIC = b"PK\x03\x04"

# Öncelik sırası önemlidir: önce tam eşleşme, sonra içerme aranır.
KEYS: dict[str, tuple[str, ...]] = {
    "date": ("işlem tarihi", "tarih", "date", "booking date", "valör"),
    "description": ("açıklama", "aciklama", "işlem açıklaması", "description", "detay", "explanation"),
    "amount": ("tutar", "işlem tutarı", "amount", "miktar"),
    "debit": ("borç", "borc", "debit", "çıkış"),
    "credit": ("alacak", "credit", "giriş"),
    "balance": ("bakiye", "kalan bakiye", "balance"),
}


class StatementError(ValueError):
    """Kullanıcıya gösterilecek kısa ayrıştırma hatası."""


# Türkçe harfler ASCII'ye iner: "İŞLEM TARİHİ", "ISLEM TARIHI", "işlem tarihi" aynı.
_ASCII = str.maketrans("çğıİIöşüÇĞÖŞÜ", "cgiiiosucgosu")


def _fold(value: object) -> str:
    # Dikkat: external_id bu normalleştirmeyle hesaplanır; değişirse
    # daha önce yüklenen ekstreler tekrar yüklenince çift kayıt oluşur.
    text = "" if value is None else str(value)
    text = text.translate(_ASCII).lower()
    return re.sub(r"\s+", " ", text).strip()


_FOLDED_KEYS: dict[str, tuple[str, ...]] = {
    field: tuple(_fold(k) for k in keys) for field, keys in KEYS.items()
}


# --- dosya okuma -----------------------------------------------------------------


def _decode(data: bytes) -> str:
    # Eski bankacılık çıktıları Windows-1254 olabilir
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1254", errors="replace")


def _read_csv(text: str) -> list[list[Cell]]:
    """Ayırıcı, başlıkta en çok anahtar kelime çıkaran adaydır.

    csv.Sniffer TR ondalık virgülünde (1.234,56) yanılabiliyor.
    """
    head = "\n".join(text.splitlines()[:HEADER_SCAN_ROWS])
    best_score, chosen = (-1, -1), ";"
    # Eşitlikte listedeki ilk aday kazanır (TR bankaları çoğunlukla ';').
    for delimiter in (";", "\t", ",", "|"):
        rows = list(csv.reader(io.StringIO(head), delimiter=delimiter))
        keys = max((sum(1 for c in r if _is_key(c)) for r in rows), default=0)
        width = max((len(r) for r in rows), default=0)
        if (keys, width) > best_score:
            best_score, chosen = (keys, width), delimiter
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=chosen)]


class _TableParser(HTMLParser):
    """'xls' adıyla inen HTML tablosu: <tr>/<td>/<th> hücreleri."""

    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list[Cell]] = []
        self._row: list[Cell] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def _read_xlsx(data: bytes) -> list[list[Cell]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl bozuk dosyada çeşitli hatalar fırlatır
        raise StatementError("Excel dosyası okunamadı") from exc
    try:
        sheet = book.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()


def read_rows(filename: str, data: bytes) -> list[list[Cell]]:
    """Dosya → ham satırlar. Biçim uzantıdan değil içerikten anlaşılır."""
    if not data:
        raise StatementError("Dosya boş")
    if len(data) > MAX_FILE_BYTES:
        raise StatementError("Dosya 3 MB'tan büyük; tarih aralığını daraltıp tekrar indirin")
    if data.startswith(ZIP_MAGIC):
        return _read_xlsx(data)
    if data.startswith(OLE2_MAGIC):
        raise StatementError("Eski .xls biçimi okunamıyor; ekstreyi XLSX veya CSV olarak indirin")
    text = _decode(data)
    if re.match(r"\s*<", text) and re.search(r"<t[rd][\s>]", text, re.I):
        parser = _TableParser()
        parser.feed(text)
        return parser.rows
    return _read_csv(text)


# --- başlık ve eşleme ------------------------------------------------------------


def _is_key(cell: object) -> bool:
    folded = _fold(cell)
    if not folded:
        return False
    return any(folded == k or folded.startswith(k) for keys in _FOLDED_KEYS.values() for k in keys)


@dataclass(frozen=True)
class Sheet:
    header: list[str]
    rows: list[list[Cell]]
    header_index: int


def find_header(raw: Sequence[Sequence[Cell]]) -> Sheet:
    """İlk 30 satırda anahtar kelimesi en çok olan satır başlıktır."""
    best, best_score = 0, -1
    for i, row in enumerate(raw[:HEADER_SCAN_ROWS]):
        score = sum(1 for cell in row if _is_key(cell))
        if score > best_score:
            best, best_score = i, score
    if not raw or best_score <= 0:
        raise StatementError("Başlık satırı bulunamadı (tarih, açıklama, tutar)")
    header = [str(c).strip() if c not in (None, "") else f"Sütun {i + 1}" for i, c in enumerate(raw[best])]
    rows = [list(r) for r in raw[best + 1 :] if any(c not in (None, "") for c in r)]
    return Sheet(header=header, rows=rows, header_index=best)


class StatementMapping(BaseModel):
    """Sütun indeksleri. amount yoksa borç/alacak sütunları kullanılır."""

    date: int
    description: int | None = None
    amount: int | None = None
    debit: int | None = None
    credit: int | None = None
    balance: int | None = None

    @model_validator(mode="after")
    def _needs_amount(self) -> "StatementMapping":
        if self.amount is None and self.debit is None and self.credit is None:
            raise ValueError("Tutar veya borç/alacak sütunu seçilmeli")
        return self


def _find(header: Sequence[str], keys: Sequence[str], taken: set[int]) -> int | None:
    folded = [_fold(h) for h in header]
    for key in keys:
        for i, h in enumerate(folded):
            if i not in taken and h == key:
                return i
    for key in keys:
        for i, h in enumerate(folded):
            if i not in taken and key in h:
                return i
    return None


def guess_mapping(header: Sequence[str]) -> StatementMapping | None:
    """Başlıktan sütun tahmini; tarih bulunamazsa None (kullanıcı seçer)."""
    taken: set[int] = set()
    found: dict[str, int | None] = {}
    # Borç/alacak önce aranır: "Borç Tutarı" tutar sütunu sanılmasın
    for field in ("date", "debit", "credit", "balance", "description", "amount"):
        index = _find(header, _FOLDED_KEYS[field], taken)
        found[field] = index
        if index is not None:
            taken.add(index)
    if found["debit"] is not None and found["credit"] is not None:
        found["amount"] = None
    elif found["amount"] is not None:
        found["debit"] = found["credit"] = None
    if found["date"] is None or (
        found["amount"] is None and found["debit"] is None and found["credit"] is None
    ):
        return None
    return StatementMapping(**found)


# --- sayı ve tarih ---------------------------------------------------------------


def parse_number(value: Cell) -> Decimal | None:
    """'1.234,56' (TR), '1,234.56' (EN), '-1234.5', '(45,10)', '45,10-' → Decimal."""
    if value is None or isinstance(value, (date, datetime)):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, Decimal)):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(repr(value))
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith("-") or text.endswith("-") or (text.startswith("(") and text.endswith(")"))
    digits = re.sub(r"[^\d,.]", "", text)
    if not re.search(r"\d", digits):
        return None
    last_comma, last_dot = digits.rfind(","), digits.rfind(".")
    if last_comma >= 0 and last_dot >= 0:
        # İkisi de varsa sondaki ondalıktır
        if last_comma > last_dot:
            digits = digits.replace(".", "").replace(",", ".")
        else:
            digits = digits.replace(",", "")
    elif last_comma >= 0 or last_dot >= 0:
        sep = "," if last_comma >= 0 else "."
        tail = len(digits) - digits.rfind(sep) - 1
        # Tek tür ayırıcı: birden çoksa veya ardında tam 3 hane varsa binliktir
        if digits.count(sep) > 1 or tail == 3:
            digits = digits.replace(sep, "")
        else:
            digits = digits.replace(sep, ".")
    try:
        number = Decimal(digits)
    except InvalidOperation:
        return None
    return -number if negative else number


def parse_date(value: Cell) -> date | None:
    """dd.mm.yyyy, dd/mm/yyyy, dd-mm-yy, yyyy-mm-dd, Excel seri sayı, datetime."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, Decimal)):
        serial = float(value)
        if not EXCEL_SERIAL_MIN <= serial <= EXCEL_SERIAL_MAX:
            return None
        return EXCEL_EPOCH + timedelta(days=int(serial))
    text = str(value).strip()
    try:
        if m := re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text):
            return date(int(m[1]), int(m[2]), int(m[3]))
        if m := re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", text):
            year = int(m[3]) + (2000 if len(m[3]) == 2 else 0)
            return date(year, int(m[2]), int(m[1]))
    except ValueError:
        return None
    if re.fullmatch(r"\d{5}(\.\d+)?", text):
        return parse_date(float(text))
    return None


# --- satırlar → hareketler -------------------------------------------------------


@dataclass(frozen=True)
class ParsedTx:
    booked_at: date
    amount: Decimal
    description: str
    balance: Decimal | None
    row: int  # başlıktan sonraki sıra (0'dan)
    external_id: str


@dataclass(frozen=True)
class Parsed:
    txs: list[ParsedTx]
    skipped: int

    @property
    def last_balance(self) -> Decimal | None:
        """En geç tarihli satırın bakiyesi.

        Ekstre eskiden yeniye veya yeniden eskiye sıralı olabilir; aynı
        gündeki satırlarda dosyanın yönüne göre en son işlem seçilir.
        """
        with_balance = [t for t in self.txs if t.balance is not None]
        if not with_balance:
            return None
        newest_first = self.txs[0].booked_at > self.txs[-1].booked_at
        latest = max(
            with_balance,
            key=lambda t: (t.booked_at, -t.row if newest_first else t.row),
        )
        return latest.balance

    @property
    def date_range(self) -> tuple[date, date] | None:
        if not self.txs:
            return None
        days = [t.booked_at for t in self.txs]
        return min(days), max(days)


def _cell(row: Sequence[Cell], index: int | None) -> Cell:
    if index is None or index >= len(row):
        return None
    return row[index]


# Türk bankası havale açıklaması: '… - Gönderen:X, Alıcı:Y, FAST Para Transferi'
# (KuveytTürk) ya da 'Açıklama= …Gönderen= X , Alıcı= Y'; yurt dışı gelen 'Amir:X <no> <no>'.
_PARTIES = re.compile(r"g[öo]nderen\s*[:=]\s*(?P<sender>.*?)\s*,\s*al[ıi]c[ıi]\s*[:=]\s*(?P<receiver>[^,]*)", re.I)
_ORDERING = re.compile(r"(?<![a-zçğıöşü])amir\s*:\s*(?P<name>.*?)(?:\s+\d[\d\s]*)?$", re.I)
COUNTERPARTY_MAX = 120


def bank_counterparty(description: str, amount: Decimal) -> str | None:
    """Açıklamadaki karşı taraf: girişte gönderen, çıkışta alıcı; yoksa None."""
    match = _PARTIES.search(description)
    if match:
        name = match.group("sender" if amount > 0 else "receiver")
    else:
        ordering = _ORDERING.search(description)
        name = ordering.group("name") if ordering and amount > 0 else ""
    name = " ".join(name.split())
    return name[:COUNTERPARTY_MAX] or None


def external_id(booked_at: date, amount: Decimal, description: str, occurrence: int) -> str:
    """İçerikten kimlik. Aynı gün, tutar, açıklama tekrarında sıra no ile ayrılır."""
    base = f"{booked_at.isoformat()}|{amount.quantize(Decimal('0.01'))}|{_fold(description)}|{occurrence}"
    return "csv:" + hashlib.sha1(base.encode("utf-8")).hexdigest()


def apply_mapping(sheet: Sheet, mapping: StatementMapping, invert: bool = False) -> Parsed:
    """Tarihi veya tutarı olmayan (ya da sıfır) satırlar atlanır: devir, ara toplam."""
    txs: list[ParsedTx] = []
    skipped = 0
    seen: dict[str, int] = {}
    for i, row in enumerate(sheet.rows):
        booked_at = parse_date(_cell(row, mapping.date))
        if mapping.amount is not None:
            amount = parse_number(_cell(row, mapping.amount))
        else:
            debit = abs(parse_number(_cell(row, mapping.debit)) or Decimal("0"))
            credit = abs(parse_number(_cell(row, mapping.credit)) or Decimal("0"))
            amount = credit - debit if (debit or credit) else None
        if booked_at is None or amount is None or amount == 0:
            skipped += 1
            continue
        if invert:
            amount = -amount
        amount = amount.quantize(Decimal("0.01"))
        raw_desc = _cell(row, mapping.description)
        description = " ".join(str(raw_desc).split()) if raw_desc is not None else ""
        key = f"{booked_at}|{amount}|{_fold(description)}"
        seen[key] = seen.get(key, 0) + 1
        balance = parse_number(_cell(row, mapping.balance))
        txs.append(
            ParsedTx(
                booked_at=booked_at,
                amount=amount,
                description=description,
                balance=balance.quantize(Decimal("0.01")) if balance is not None else None,
                row=i,
                external_id=external_id(booked_at, amount, description, seen[key]),
            )
        )
    return Parsed(txs=txs, skipped=skipped)
